"""FAQ 知識庫維護：讀寫 data/faq.csv、匯入試算表，並即時同步 Qdrant 向量庫。"""

import csv
import io
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import requests

from .config import BATCH_SIZE, PROJECT_ROOT
from .embedding import embed_texts
from .qdrant_service import delete_faqs, recreate_collection, upsert_faqs

CSV_PATH = PROJECT_ROOT / "data" / "faq.csv"
BACKUP_DIR = PROJECT_ROOT / "data" / "faq-backups"
BACKUP_KEEP = 20
FIELDS = ["id", "category", "question", "answer", "url", "keywords", "office", "email"]
TAIPEI = timezone(timedelta(hours=8))
MAX_ROWS = 5000

# 試算表標題（中文或英文）對應到 FAQ 欄位，比對前會先去空白並轉小寫。
HEADER_ALIASES = {
    "id": "id", "問題編號": "id", "編號": "id", "序號": "id", "題號": "id", "faq id": "id", "faq編號": "id",
    "category": "category", "分類": "category", "類別": "category", "主題": "category", "主題分類": "category",
    "question": "question", "問題": "question", "常見問題": "question", "題目": "question", "提問": "question",
    "answer": "answer", "答案": "answer", "建議答案": "answer", "回答": "answer", "參考答案": "answer", "說明": "answer",
    "url": "url", "網址": "url", "連結": "url", "來源": "url", "來源連結": "url", "依據": "url",
    "依據/來源連結": "url", "依據及來源連結": "url", "參考連結": "url", "資料來源": "url",
    "keywords": "keywords", "關鍵字": "keywords", "標籤": "keywords", "同義詞": "keywords",
    "office": "office", "主責單位": "office", "承辦單位": "office", "負責單位": "office", "處室": "office",
    "單位": "office", "權責單位": "office",
    "email": "email", "信箱": "email", "電子信箱": "email", "聯絡信箱": "email", "承辦信箱": "email",
    "e-mail": "email", "mail": "email",
}
_ENCODINGS = ("utf-8-sig", "utf-8", "cp950", "big5")


def _norm_header(value: Any) -> str:
    text = str(value or "").replace("﻿", "").replace("／", "/").replace("　", " ").strip().lower()
    return re.sub(r"\s+", "", text)


_ALIASES = {_norm_header(key): field for key, field in HEADER_ALIASES.items()}


def _clean(value: Any) -> str:
    text = "" if value is None else str(value)
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def today() -> str:
    return datetime.now(TAIPEI).strftime("%Y-%m-%d")


# ---------------------------------------------------------------- CSV 讀寫

def read_rows() -> list[dict[str, str]]:
    """讀取 faq.csv，缺少的欄位一律補空字串。"""
    if not CSV_PATH.exists():
        return []
    last_error: Exception | None = None
    for encoding in _ENCODINGS:
        try:
            with CSV_PATH.open("r", encoding=encoding, newline="") as source:
                return [{field: _clean(row.get(field)) for field in FIELDS} for row in csv.DictReader(source)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"faq.csv 不是支援的編碼（UTF-8 或 Big5）：{last_error}")


def write_rows(rows: list[dict[str, str]]) -> None:
    """以暫存檔改名的方式寫回 faq.csv，避免寫到一半損毀原檔。"""
    CSV_PATH.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(str(CSV_PATH) + ".tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows([{field: row.get(field, "") for field in FIELDS} for row in rows])
    temporary.replace(CSV_PATH)


def backup_csv() -> str:
    """匯入或大量異動前先備份，只保留最近幾份。"""
    if not CSV_PATH.exists():
        return ""
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    name = f"faq-{datetime.now(TAIPEI).strftime('%Y%m%d-%H%M%S')}.csv"
    (BACKUP_DIR / name).write_bytes(CSV_PATH.read_bytes())
    for stale in sorted(BACKUP_DIR.glob("faq-*.csv"))[:-BACKUP_KEEP]:
        stale.unlink(missing_ok=True)
    return name


# ---------------------------------------------------------------- 向量同步

def embedding_text(row: dict[str, str]) -> str:
    keywords = " ".join(part for part in (row.get("keywords", ""), row.get("office", "")) if part).strip()
    return (f"分類：\n{row.get('category', '')}\n\n問題：\n{row['question']}\n\n"
            f"答案：\n{row['answer']}\n\n關鍵字：\n{keywords}")


def _payload(row: dict[str, str]) -> dict[str, str]:
    return {**{field: row.get(field, "") for field in FIELDS}, "updated_at": row.get("updated_at") or today()}


def sync_vectors(rows: list[dict[str, str]]) -> int:
    """把 FAQ 逐批做 embedding 後寫入 Qdrant，回傳實際寫入筆數。"""
    usable = [row for row in rows if row.get("question") and row.get("answer")]
    for start in range(0, len(usable), BATCH_SIZE):
        chunk = usable[start:start + BATCH_SIZE]
        upsert_faqs([_payload(row) for row in chunk], embed_texts([embedding_text(row) for row in chunk]))
    return len(usable)


def reindex() -> int:
    """重建整個向量庫，用於匯入異常或手動修復。"""
    rows = read_rows()
    recreate_collection()
    return sync_vectors(rows)


# ---------------------------------------------------------------- 單筆維護

def _next_id(used: set[str]) -> str:
    numbers = [int(match.group(1)) for value in used if (match := re.fullmatch(r"[Qq](\d+)", value))]
    return f"Q{(max(numbers) + 1) if numbers else 1:03d}"


def normalize_id(raw: Any, used: set[str] | None = None) -> str:
    """試算表的「問題編號」多半是 1、2、3，統一轉成 Q001 這種固定 ID。"""
    value = _clean(raw)
    if value.isdigit():
        return f"Q{int(value):03d}"
    if value:
        return value
    return _next_id(used or set())


def list_faqs() -> list[dict[str, str]]:
    return read_rows()


def summary() -> dict[str, Any]:
    rows = read_rows()
    categories: dict[str, int] = {}
    offices: dict[str, int] = {}
    for row in rows:
        name = row["category"] or "未分類"
        categories[name] = categories.get(name, 0) + 1
        if row["office"]:
            offices[row["office"]] = offices.get(row["office"], 0) + 1
    return {
        "total": len(rows),
        "categories": sorted(({"name": name, "total": total} for name, total in categories.items()),
                             key=lambda item: -item["total"]),
        "offices": sorted(({"name": name, "total": total} for name, total in offices.items()),
                          key=lambda item: -item["total"]),
        "updated_at": datetime.fromtimestamp(CSV_PATH.stat().st_mtime, TAIPEI).strftime("%Y-%m-%d %H:%M")
        if CSV_PATH.exists() else "",
    }


def save_faq(record: dict[str, Any], original_id: str = "") -> dict[str, str]:
    """新增或更新單筆 FAQ，CSV 與向量庫一起更新。"""
    rows = read_rows()
    used = {row["id"] for row in rows}
    row = {field: _clean(record.get(field)) for field in FIELDS}
    if not row["question"] or not row["answer"]:
        raise ValueError("問題與答案不可為空白")
    original_id = _clean(original_id)
    row["id"] = normalize_id(row["id"] or original_id, used)
    if row["id"] != original_id and row["id"] in used:
        raise ValueError(f"問題編號 {row['id']} 已存在，請換一個編號")
    if not row["keywords"]:
        row["keywords"] = " ".join(part for part in (row["category"], row["office"]) if part)
    target = original_id or row["id"]
    for index, existing in enumerate(rows):
        if existing["id"] == target:
            rows[index] = row
            break
    else:
        rows.append(row)
    write_rows(rows)
    sync_vectors([row])
    if original_id and original_id != row["id"]:
        delete_faqs([original_id])
    return row


def delete_faq(faq_id: str) -> dict[str, str]:
    faq_id = _clean(faq_id)
    rows = read_rows()
    kept = [row for row in rows if row["id"] != faq_id]
    if len(kept) == len(rows):
        raise KeyError(f"找不到問題編號 {faq_id}")
    write_rows(kept)
    delete_faqs([faq_id])
    return {"id": faq_id}


# ---------------------------------------------------------------- 匯入來源解析

def _decode(raw: bytes) -> str:
    for encoding in _ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("檔案不是支援的編碼，請另存為 UTF-8 CSV 後再上傳")


def _table_from_text(text: str) -> list[list[str]]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip("\n")
    if not text.strip():
        raise ValueError("沒有讀到任何內容")
    head = text.split("\n", 1)[0]
    delimiter = "\t" if "\t" in head and head.count("\t") >= head.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter) if any(_clean(cell) for cell in row)]


def _table_from_xlsx(raw: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # 只有伺服器未安裝 openpyxl 時才會發生
        raise ValueError("伺服器缺少 openpyxl，無法讀取 .xlsx，請改上傳 CSV") from exc
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    table = [[_clean(cell) for cell in row] for row in workbook.active.iter_rows(values_only=True)]
    workbook.close()
    return [row for row in table if any(row)]


def sheet_csv_url(url: str) -> str:
    """把 Google 試算表的編輯網址轉成 CSV 匯出網址；其他網址原樣回傳。"""
    url = _clean(url)
    match = re.search(r"/spreadsheets/d/([A-Za-z0-9\-_]+)", url)
    if not match:
        return url
    gid = re.search(r"[#&?]gid=(\d+)", url)
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid.group(1) if gid else '0'}"


def fetch_source(url: str) -> list[list[str]]:
    target = sheet_csv_url(url)
    if not target.startswith(("http://", "https://")):
        raise ValueError("請輸入完整的 http/https 網址")
    try:
        response = requests.get(target, timeout=30, allow_redirects=True)
    except requests.RequestException as exc:
        raise ValueError(f"無法連線到試算表：{exc}") from exc
    if response.status_code != 200:
        raise ValueError(f"下載試算表失敗（HTTP {response.status_code}），請確認已開放「知道連結的人可以檢視」")
    if "html" in response.headers.get("Content-Type", "").lower():
        raise ValueError("試算表沒有開放檢視權限，請改成「知道連結的人可以檢視」，或直接上傳檔案")
    return _table_from_text(_decode(response.content))


def parse_text(text: str) -> list[list[str]]:
    """使用者直接從試算表複製貼上的內容（TSV 或 CSV）。"""
    return _table_from_text(text)


def parse_upload(raw: bytes, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _table_from_xlsx(raw)
    if name.endswith(".xls"):
        raise ValueError("不支援舊版 .xls，請另存為 .xlsx 或 CSV")
    return _table_from_text(_decode(raw))


# ---------------------------------------------------------------- 匯入預覽與寫入

def map_table(table: list[list[str]]) -> tuple[list[dict[str, str]], dict[str, str]]:
    """找出標題列並把每列轉成 FAQ 欄位，回傳（資料列, 欄位對應）。"""
    for index, row in enumerate(table[:10]):
        mapping = {position: _ALIASES[key] for position, cell in enumerate(row)
                   if (key := _norm_header(cell)) in _ALIASES}
        if "question" in mapping.values() and "answer" in mapping.values():
            header = {mapping[position]: _clean(row[position]) for position in mapping}
            records = []
            for line in table[index + 1:]:
                record = {field: _clean(line[position]) if position < len(line) else ""
                          for position, field in mapping.items()}
                if any(record.values()):
                    records.append(record)
            return records, header
    raise ValueError("找不到標題列，請確認第一列含有「問題」與「建議答案」欄位")


def plan_import(records: list[dict[str, str]]) -> dict[str, Any]:
    """比對現有 FAQ，標出每一列會新增、更新，還是因為缺欄位被略過。"""
    if len(records) > MAX_ROWS:
        raise ValueError(f"一次最多匯入 {MAX_ROWS} 筆，目前有 {len(records)} 筆")
    existing = {row["id"]: row for row in read_rows()}
    used = set(existing)
    rows, counts = [], {"create": 0, "update": 0, "skip": 0}
    for number, record in enumerate(records, start=2):
        row = {field: _clean(record.get(field)) for field in FIELDS}
        missing = [name for field, name in (("question", "問題"), ("answer", "建議答案")) if not row[field]]
        if missing:
            rows.append({**row, "row": number, "action": "skip", "note": "缺少" + "、".join(missing)})
            counts["skip"] += 1
            continue
        row["id"] = normalize_id(row["id"], used)
        used.add(row["id"])
        if not row["keywords"]:
            row["keywords"] = " ".join(part for part in (row["category"], row["office"]) if part)
        action = "update" if row["id"] in existing else "create"
        note = f"覆蓋原本的「{existing[row['id']]['question'][:20]}」" if action == "update" else "新增"
        rows.append({**row, "row": number, "action": action, "note": note})
        counts[action] += 1
    return {"rows": rows, "counts": counts, "existing_total": len(existing)}


def apply_import(rows: list[dict[str, Any]], mode: str = "merge") -> dict[str, Any]:
    """把預覽通過的資料寫入 CSV 與向量庫。merge 是合併更新，replace 是整份取代。"""
    if mode not in ("merge", "replace"):
        raise ValueError("匯入模式只能是 merge 或 replace")
    records = [{field: _clean(row.get(field)) for field in FIELDS} for row in rows
               if _clean(row.get("question")) and _clean(row.get("answer"))]
    if not records:
        raise ValueError("沒有可匯入的資料")
    used: set[str] = set()
    for record in records:
        record["id"] = normalize_id(record["id"], used)
        used.add(record["id"])
    backup = backup_csv()
    if mode == "replace":
        recreate_collection()
        write_rows(records)
        sync_vectors(records)
        return {"mode": mode, "created": len(records), "updated": 0, "total": len(records), "backup": backup}
    existing = read_rows()
    index = {row["id"]: position for position, row in enumerate(existing)}
    created = updated = 0
    for record in records:
        if record["id"] in index:
            existing[index[record["id"]]] = record
            updated += 1
        else:
            index[record["id"]] = len(existing)
            existing.append(record)
            created += 1
    write_rows(existing)
    sync_vectors(records)
    return {"mode": mode, "created": created, "updated": updated, "total": len(existing), "backup": backup}


def export_csv() -> str:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=FIELDS, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(read_rows())
    return "﻿" + output.getvalue()


def template_csv() -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["問題編號", "分類", "問題", "建議答案", "依據／來源連結", "主責單位", "信箱"])
    writer.writerow(["1", "選課相關", "如何加退選課程？", "請登入選課系統辦理加退選。",
                     "https://example.edu.tw/course", "教務處", "example@mail.shu.edu.tw"])
    return "﻿" + output.getvalue()
